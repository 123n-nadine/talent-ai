"""Stage 3: text extraction with layout.

Text PDF -> PyMuPDF (bidi row reconstruction), DOCX -> python-docx (+ a LibreOffice PDF preview
for the page view), DOC/RTF -> LibreOffice PDF, scanned PDF or image -> OCR provider.

Besides the text, every extraction returns a *layout*: page sizes and one entry per text line with
its box on the page and its confidence (1.0 for a text layer, the recogniser score for OCR), plus
the OCR word/segment boxes. The CV analysis draws these over the original page. Coordinates are in
the pixel space of the page image the engine serves (`DISPLAY_DPI` for PDFs, native pixels for
images). Any failure returns status 'unreadable' with a reason; the screening then continues from
the form answers (decision 14)."""

from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from app.core.config import Settings
from app.documents.pdftext import group_rows, page_spans, row_text
from app.ocr.base import OCRPage, OCRProvider, OCRUnavailable

LETTERS = re.compile(r"[\w؀-ۿ]", re.UNICODE)
DISPLAY_DPI = 110  # page images served to the UI are rendered at this resolution


@dataclass
class Extraction:
    status: str  # ready | unreadable
    text: str = ""
    method: str = ""
    pages: int = 0
    chars: int = 0
    ocr_confidence: float | None = None
    error: str | None = None
    layout: dict[str, Any] | None = None
    preview_pdf: bytes | None = None  # rendered pages for DOCX/DOC/RTF


def _chars(text: str) -> int:
    return len(LETTERS.findall(text))


def _norm_line(line: str) -> str:
    line = line.replace("\x00", "")
    return re.sub(r"[ \t ]+", " ", line).strip()


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [_norm_line(line) for line in text.split("\n")]
    out = "\n".join(lines)
    return re.sub(r"\n{3,}", "\n\n", out).strip()


def _new_layout(source: str, dpi: int | None) -> dict[str, Any]:
    return {"source": source, "dpi": dpi, "pages": [], "lines": []}


def _pdf_text_layout(doc: Any, source: str) -> tuple[list[str], dict[str, Any]]:
    """Lines and layout from a PDF text layer (one box per reconstructed row)."""
    layout = _new_layout(source, DISPLAY_DPI)
    k = DISPLAY_DPI / 72.0
    lines: list[str] = []
    for pno in range(doc.page_count):
        page = doc.load_page(pno)
        layout["pages"].append(
            {"w": round(page.rect.width * k, 1), "h": round(page.rect.height * k, 1)}
        )
        for row in group_rows(page_spans(page)):
            t = _norm_line(row_text(row))
            if not t:
                continue
            x0 = min(s["x0"] for s in row) * k
            x1 = max(s["x1"] for s in row) * k
            y0 = min(s["yc"] - s["h"] / 2 for s in row) * k
            y1 = max(s["yc"] + s["h"] / 2 for s in row) * k
            box = [round(x0, 1), round(y0, 1), round(x1, 1), round(y1, 1)]
            layout["lines"].append(
                {"p": pno, "t": t, "c": 1.0, "b": box, "boxes": [box + [1.0, t]]}
            )
            lines.append(t)
    return lines, layout


def _ocr_layout(
    pages: list[OCRPage],
    source: str,
    dpi: int | None,
    scale: float = 1.0,
    page_sizes: list[tuple[float, float]] | None = None,
) -> tuple[list[str], dict[str, Any]]:
    layout = _new_layout(source, dpi)
    lines: list[str] = []
    for pno, pg in enumerate(pages):
        w, h = (
            page_sizes[pno]
            if page_sizes and pno < len(page_sizes)
            else (pg.width * scale, pg.height * scale)
        )
        layout["pages"].append({"w": round(w, 1), "h": round(h, 1)})
        for ln in pg.lines or []:
            t = _norm_line(ln.text)
            if not t:
                continue
            boxes = [
                [
                    round(b.x0 * scale, 1),
                    round(b.y0 * scale, 1),
                    round(b.x1 * scale, 1),
                    round(b.y1 * scale, 1),
                    round(b.confidence, 4),
                    b.text,
                ]
                for b in ln.boxes
            ]
            if boxes:
                bb = [
                    min(float(b.x0) for b in ln.boxes) * scale,
                    min(float(b.y0) for b in ln.boxes) * scale,
                    max(float(b.x1) for b in ln.boxes) * scale,
                    max(float(b.y1) for b in ln.boxes) * scale,
                ]
                bb = [round(v, 1) for v in bb]
            else:
                bb = [0.0, 0.0, 0.0, 0.0]
            layout["lines"].append(
                {"p": pno, "t": t, "c": round(ln.confidence, 4), "b": bb, "boxes": boxes}
            )
            lines.append(t)
        if not pg.lines and pg.text.strip():  # provider without line detail
            for t in (_norm_line(x) for x in pg.text.split("\n")):
                if t:
                    layout["lines"].append(
                        {
                            "p": pno,
                            "t": t,
                            "c": round(pg.confidence, 4),
                            "b": [0, 0, 0, 0],
                            "boxes": [],
                        }
                    )
                    lines.append(t)
    return lines, layout


def _run_ocr(paths: list[Path], ocr: OCRProvider | None) -> tuple[list[OCRPage], float]:
    if ocr is None:
        raise OCRUnavailable("no OCR provider configured (ocr_provider = none)")
    pages: list[OCRPage] = []
    for p in paths:
        pages.extend(ocr.extract(str(p)))
    confs = [pg.confidence for pg in pages]
    return pages, (round(sum(confs) / len(confs), 4) if confs else 0.0)


def extract_pdf(
    path: Path, ocr: OCRProvider | None, settings: Settings, provider_name: str, source: str = "pdf"
) -> Extraction:
    import pymupdf

    try:
        doc = pymupdf.open(path)
    except Exception as e:  # noqa: BLE001
        msg = str(e).replace(str(path), "<file>")  # never expose the storage path
        return Extraction("unreadable", error=f"corrupt PDF: {type(e).__name__}: {msg}")
    with doc:
        if doc.needs_pass:
            return Extraction(
                "unreadable", error="encrypted PDF (password protected)", pages=doc.page_count
            )
        if doc.page_count == 0:
            return Extraction("unreadable", error="PDF has no pages")
        try:
            lines, layout = _pdf_text_layout(doc, source)
        except Exception as e:  # noqa: BLE001
            return Extraction("unreadable", error=f"corrupt PDF content: {type(e).__name__}: {e}")
        text = "\n".join(lines)
        if _chars(text) >= settings.pdf_min_text_chars:
            return Extraction(
                "ready",
                text=text,
                method="pdf_text",
                pages=doc.page_count,
                chars=_chars(text),
                layout=layout,
            )
        # no usable text layer -> scanned -> OCR each rendered page
        sizes = [
            (
                doc.load_page(i).rect.width * DISPLAY_DPI / 72,
                doc.load_page(i).rect.height * DISPLAY_DPI / 72,
            )
            for i in range(doc.page_count)
        ]
        with tempfile.TemporaryDirectory(prefix="talentai-ocr-") as tmp:
            images: list[Path] = []
            try:
                for i in range(doc.page_count):
                    pix = doc.load_page(i).get_pixmap(dpi=settings.ocr_render_dpi)
                    out = Path(tmp) / f"page-{i + 1}.png"
                    pix.save(out)
                    images.append(out)
                ocr_pages, conf = _run_ocr(images, ocr)
            except OCRUnavailable as e:
                return Extraction(
                    "unreadable",
                    error=f"scanned PDF but OCR is unavailable: {e}",
                    pages=doc.page_count,
                )
            except Exception as e:  # noqa: BLE001
                return Extraction(
                    "unreadable", error=f"OCR failed: {type(e).__name__}: {e}", pages=doc.page_count
                )
        lines, layout = _ocr_layout(
            ocr_pages,
            "pdf",
            DISPLAY_DPI,
            scale=DISPLAY_DPI / settings.ocr_render_dpi,
            page_sizes=sizes,
        )
        ocr_text = "\n".join(lines)
        if not _chars(ocr_text):
            return Extraction(
                "unreadable",
                error="OCR produced no text",
                pages=doc.page_count,
                method=f"ocr:{provider_name}",
                ocr_confidence=conf,
                layout=layout,
            )
        return Extraction(
            "ready",
            text=ocr_text,
            method=f"ocr:{provider_name}",
            pages=doc.page_count,
            chars=_chars(ocr_text),
            ocr_confidence=conf,
            layout=layout,
        )


def extract_docx(path: Path) -> Extraction:
    import docx

    try:
        d = docx.Document(str(path))
    except Exception as e:  # noqa: BLE001
        return Extraction("unreadable", error=f"corrupt DOCX: {type(e).__name__}: {e}")
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = "\n".join(t for t in (_norm_line(x) for x in parts) if t)
    if not _chars(text):
        return Extraction("unreadable", error="DOCX contains no text", method="docx")
    return Extraction("ready", text=text, method="docx", pages=1, chars=_chars(text))


def convert_with_libreoffice(path: Path, settings: Settings, outdir: Path) -> Path:
    exe = shutil.which(settings.libreoffice_bin)
    if not exe:
        raise RuntimeError("LibreOffice (soffice) is not installed")
    profile = outdir / "lo-profile"
    cmd = [
        exe,
        "--headless",
        "--norestore",
        "--nologo",
        f"-env:UserInstallation=file://{profile}",
        "--convert-to",
        "pdf",
        "--outdir",
        str(outdir),
        str(path),
    ]
    subprocess.run(cmd, check=True, capture_output=True, timeout=settings.libreoffice_timeout_s)
    pdf = outdir / (path.stem + ".pdf")
    if not pdf.exists():
        raise RuntimeError("LibreOffice produced no PDF")
    return pdf


def extract_image(path: Path, ocr: OCRProvider | None, provider_name: str) -> Extraction:
    try:
        with tempfile.TemporaryDirectory(prefix="talentai-img-") as tmp:
            # stored files have no extension; OCR engines want a real image suffix
            head = path.read_bytes()[:8]
            suffix = ".png" if head.startswith(b"\x89PNG") else ".jpg"
            img = Path(tmp) / f"scan{suffix}"
            shutil.copyfile(path, img)
            ocr_pages, conf = _run_ocr([img], ocr)
    except OCRUnavailable as e:
        return Extraction("unreadable", error=f"image CV but OCR is unavailable: {e}", pages=1)
    except Exception as e:  # noqa: BLE001
        return Extraction("unreadable", error=f"OCR failed: {type(e).__name__}: {e}", pages=1)
    lines, layout = _ocr_layout(ocr_pages, "image", None)
    text = "\n".join(lines)
    if not _chars(text):
        return Extraction(
            "unreadable",
            error="OCR produced no text",
            pages=1,
            method=f"ocr:{provider_name}",
            ocr_confidence=conf,
            layout=layout,
        )
    return Extraction(
        "ready",
        text=text,
        method=f"ocr:{provider_name}",
        pages=1,
        chars=_chars(text),
        ocr_confidence=conf,
        layout=layout,
    )


def _preview_layout(pdf: Path) -> tuple[dict[str, Any] | None, int]:
    import pymupdf

    try:
        with pymupdf.open(pdf) as doc:
            _, layout = _pdf_text_layout(doc, "preview_pdf")
            return layout, doc.page_count
    except Exception:  # noqa: BLE001
        return None, 0


def extract(
    path: Path,
    kind: str,
    ocr: OCRProvider | None,
    settings: Settings,
    provider_name: str = "none",
    preview: bool = True,
) -> Extraction:
    if kind == "pdf":
        return extract_pdf(path, ocr, settings, provider_name)
    if kind == "docx":
        res = extract_docx(path)
        if res.status == "ready" and preview:
            with tempfile.TemporaryDirectory(prefix="talentai-lo-") as tmp:
                try:
                    pdf = convert_with_libreoffice(path, settings, Path(tmp))
                    res.layout, pages = _preview_layout(pdf)
                    res.pages = pages or res.pages
                    res.preview_pdf = pdf.read_bytes()
                except Exception:  # noqa: BLE001 - the preview is optional for DOCX
                    pass
        return res
    if kind in ("doc", "rtf"):
        with tempfile.TemporaryDirectory(prefix="talentai-lo-") as tmp:
            try:
                pdf = convert_with_libreoffice(path, settings, Path(tmp))
            except Exception as e:  # noqa: BLE001
                return Extraction("unreadable", error=f"{kind.upper()} conversion failed: {e}")
            res = extract_pdf(pdf, ocr, settings, provider_name, source="preview_pdf")
            if res.status == "ready":
                res.method = f"libreoffice+{res.method}"
                res.preview_pdf = pdf.read_bytes()
            return res
    if kind == "image":
        return extract_image(path, ocr, provider_name)
    return Extraction("unreadable", error=f"unsupported file type '{kind}'")


__all__ = ["DISPLAY_DPI", "Extraction", "extract", "convert_with_libreoffice", "_clean"]

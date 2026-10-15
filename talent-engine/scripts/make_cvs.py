"""Generate ~20 synthetic CV files into data/cvs/ (git-ignored) from scripts/cv_texts.py.

Types: Arabic text PDF, English text PDF, DOCX, scanned Arabic image (PDF rendered to JPEG with
noise and a slight rotation), one corrupt PDF and one encrypted PDF. Fictional data only.
Needs LibreOffice (soffice) for DOCX -> PDF with Arabic shaping. Usage: make cvs
"""

from __future__ import annotations

import os
import random
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from scripts.cv_texts import CV_TEXTS  # noqa: E402
from scripts.seed_data import CANDIDATES  # noqa: E402

OUT = Path(os.environ.get("CV_OUT_DIR", Path(__file__).resolve().parent.parent / "data" / "cvs"))
FEMALE = {"HR-APP-2026-1951", "HR-APP-2026-1952", "HR-APP-2026-1953"}
FACTS = {c["ref"]: c["facts"] for c in CANDIDATES}


def is_arabic(s: str) -> bool:
    return bool(re.search(r"[؀-ۿ]", s))


def contact_block(ref: str, first_line: str) -> tuple[str, list[str]]:
    """Fictional personal block the redactor must remove. Returns (name, lines)."""
    parts = [x.strip() for x in first_line.replace(" - ", " — ").split(" — ") if x.strip()]
    generic = {"السيرة الذاتية", "سيرة ذاتية", "cv", "resume", "curriculum vitae"}
    name = next((x for x in parts if x.lower() not in generic), parts[0])
    digits = re.sub(r"\D", "", ref)[-4:]
    facts = FACTS.get(ref, {})
    age = int(facts.get("age") or 30)
    year = 2026 - age
    nat = facts.get("nationality") or "السعودية"
    ar = is_arabic(first_line)
    gender = ("أنثى" if ar else "Female") if ref in FEMALE else ("ذكر" if ar else "Male")
    if ar:
        lines = [
            f"الاسم: {name}",
            f"الجوال: 05{digits}{digits[::-1]}",
            f"البريد الإلكتروني: demo.{digits}@example.com",
            f"رقم الهوية الوطنية: 1{digits}{digits}{digits[:1]}",
            f"تاريخ الميلاد: 14/05/{year}",
            f"الجنسية: {nat}",
            f"الجنس: {gender}",
            "العنوان: حي النخيل، الرياض",
        ]
    else:
        lines = [
            f"Name: {name}",
            f"Mobile: +966 5{digits} {digits[::-1]}",
            f"Email: demo.{digits}@example.com",
            f"National ID: 1{digits}{digits}{digits[:1]}",
            f"Date of birth: 14/05/{year}",
            f"Nationality: {'Saudi' if nat == 'السعودية' else nat}",
            f"Gender: {gender}",
            "Address: Al Nakheel district, Riyadh",
            f"linkedin.com/in/demo-{digits}",
        ]
    return name, lines


def build_docx(path: Path, title: str, lines: list[str]) -> None:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    ar = is_arabic(title)
    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Noto Naskh Arabic" if ar else "DejaVu Sans"
    style.font.size = Pt(12)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:cs"), style.font.name)
    rfonts.set(qn("w:ascii"), style.font.name)
    rfonts.set(qn("w:hAnsi"), style.font.name)

    def para(text: str, bold: bool = False, size: int = 12) -> None:
        p = d.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        if ar:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            ppr = p._p.get_or_add_pPr()
            bidi = OxmlElement("w:bidi")
            ppr.append(bidi)
            rtl = OxmlElement("w:rtl")
            run._r.get_or_add_rPr().append(rtl)

    para(title, bold=True, size=16)
    for line in lines:
        para(line)
    d.save(str(path))


def docx_to_pdf(docx_path: Path, outdir: Path) -> Path:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise SystemExit("LibreOffice (soffice) is required to make PDFs")
    with tempfile.TemporaryDirectory(prefix="lo-profile-") as prof:
        subprocess.run(
            [
                exe,
                "--headless",
                "--norestore",
                "--nologo",
                f"-env:UserInstallation=file://{prof}",
                "--convert-to",
                "pdf",
                "--outdir",
                str(outdir),
                str(docx_path),
            ],
            check=True,
            capture_output=True,
            timeout=180,
        )
    pdf = outdir / (docx_path.stem + ".pdf")
    if not pdf.exists():
        raise SystemExit(f"conversion produced no PDF for {docx_path}")
    return pdf


def pdf_to_scan(pdf_path: Path, jpg_path: Path, seed: int) -> None:
    import pymupdf
    from PIL import Image, ImageFilter

    rng = random.Random(seed)
    with pymupdf.open(pdf_path) as doc:
        page = doc[0]
        pix = page.get_pixmap(dpi=200)  # a phone/scanner image of a printed page
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    img = img.rotate(rng.uniform(-1.2, 1.2), expand=True, fillcolor=(245, 245, 240))
    img = img.filter(ImageFilter.GaussianBlur(radius=0.4))
    px = img.load()
    w, h = img.size
    for _ in range(int(w * h * 0.002)):  # speckle noise
        x, y = rng.randrange(w), rng.randrange(h)
        v = rng.choice((40, 60, 200, 230))
        px[x, y] = (v, v, v)
    img = img.convert("L").convert("RGB")
    img.save(jpg_path, "JPEG", quality=75)


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    made = []
    with tempfile.TemporaryDirectory(prefix="talentai-cvs-") as tmp:
        work = Path(tmp)
        for i, (ref, entry) in enumerate(CV_TEXTS.items()):
            cvt, lines = entry["type"], entry["lines"]
            name, contact = contact_block(ref, lines[0])
            title = lines[0]
            body = contact + [""] + lines[1:]
            docx_path = work / f"{ref}.docx"
            build_docx(docx_path, title, body)
            if ref == "HR-APP-2026-2013":
                continue
            if cvt == "docx":
                dest = OUT / f"{ref}.docx"
                shutil.copy(docx_path, dest)
            elif cvt in ("pdf_ar", "pdf_en"):
                pdf = docx_to_pdf(docx_path, work)
                dest = OUT / f"{ref}.pdf"
                shutil.copy(pdf, dest)
            elif cvt == "scan_ar":
                pdf = docx_to_pdf(docx_path, work)
                dest = OUT / f"{ref}.jpg"
                pdf_to_scan(pdf, dest, seed=i)
            else:
                continue
            made.append((ref, dest.name, name))
        # corrupt PDF for the demo's "Failed" applicant (decision 14: assessed from answers, flag cv_unreadable)
        corrupt = OUT / "HR-APP-2026-2013.pdf"
        corrupt.write_bytes(b"%PDF-1.4\n" + random.Random(13).randbytes(4096))
        made.append(("HR-APP-2026-2013", corrupt.name, "فيصل العمري"))
        # encrypted PDF sample (not attached to a candidate; for /v1/documents/inspect)
        import pymupdf

        src = docx_to_pdf(work / "HR-APP-2026-2001.docx", work)
        with pymupdf.open(src) as doc:
            doc.save(
                str(OUT / "sample-encrypted.pdf"),
                encryption=pymupdf.PDF_ENCRYPT_AES_256,
                owner_pw="owner",
                user_pw="secret",
            )
        made.append(("-", "sample-encrypted.pdf", "-"))
    names = OUT / "name_hints.txt"
    names.write_text("\n".join(f"{ref}\t{fn}\t{name}" for ref, fn, name in made) + "\n")
    for ref, fn, name in made:
        print(f"{ref:<18} {fn:<28} hint={name}")
    print(f"\n{len(made)} files in {OUT} (git-ignored). Name hints in {names.name} (fictional).")
    return 0


if __name__ == "__main__":
    sys.exit(main())

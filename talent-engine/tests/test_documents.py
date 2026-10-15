from __future__ import annotations

import io
from pathlib import Path

import pymupdf
import pytest
from PIL import Image, ImageDraw

from app.documents.detect import sniff
from app.documents.extract import extract
from app.ocr.base import NoneOCR
from app.ocr.fake import FakeOCR
from tests.conftest import make_settings


def text_pdf_bytes(lines: list[str]) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    y = 72
    for line in lines:
        page.insert_text((72, y), line, fontsize=11)
        y += 16
    return doc.tobytes()


def scanned_pdf_bytes() -> bytes:
    img = Image.new("RGB", (800, 400), "white")
    ImageDraw.Draw(img).text((20, 20), "Security guard since 2019 - scanned", fill="black")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    doc = pymupdf.open()
    page = doc.new_page(width=800, height=400)
    page.insert_image(page.rect, stream=buf.getvalue())
    return doc.tobytes()


def docx_bytes(paragraphs: list[str]) -> bytes:
    import docx

    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "رخصة"
    t.rows[0].cells[1].text = "سارية"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def png_bytes() -> bytes:
    img = Image.new("RGB", (300, 100), "white")
    ImageDraw.Draw(img).text((10, 10), "first aid", fill="black")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def test_sniff_by_magic_bytes():
    assert sniff("x.bin", text_pdf_bytes(["a"]))[0] == "pdf"
    assert sniff("cv.docx", docx_bytes(["a"]))[0] == "docx"
    assert sniff("photo", png_bytes()) == ("image", "image/png")
    assert sniff("cv.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 16)[0] == "doc"
    assert sniff("cv.rtf", b"{\\rtf1\\ansi hello}")[0] == "rtf"
    assert sniff("cv.exe", b"MZ\x90\x00")[0] == "unknown"


def test_text_pdf(tmp_path: Path):
    p = tmp_path / "cv.pdf"
    p.write_bytes(
        text_pdf_bytes(
            [
                "Security supervisor 2017-2024",
                "Valid security guard licence",
                "First aid certificate 2025",
            ]
        )
    )
    ex = extract(p, "pdf", NoneOCR(), make_settings(), "none")
    assert ex.status == "ready" and ex.method == "pdf_text" and ex.pages == 1
    assert "licence" in ex.text and ex.ocr_confidence is None and ex.chars > 50


def test_scanned_pdf_needs_ocr(tmp_path: Path):
    p = tmp_path / "scan.pdf"
    p.write_bytes(scanned_pdf_bytes())
    ex = extract(p, "pdf", NoneOCR(), make_settings(), "none")
    assert ex.status == "unreadable" and "OCR" in (ex.error or "")
    ex = extract(p, "pdf", FakeOCR("حارس أمن منذ 2019", 0.91), make_settings(), "fake")
    assert (
        ex.status == "ready"
        and ex.method == "ocr:fake"
        and ex.ocr_confidence == 0.91
        and "2019" in ex.text
    )


def test_image_cv(tmp_path: Path):
    p = tmp_path / "cv.png"
    p.write_bytes(png_bytes())
    ex = extract(p, "image", FakeOCR("إسعافات أولية", 0.7), make_settings(), "fake")
    assert ex.status == "ready" and ex.method == "ocr:fake" and ex.ocr_confidence == 0.7
    ex = extract(p, "image", NoneOCR(), make_settings(), "none")
    assert ex.status == "unreadable"


def test_docx(tmp_path: Path):
    p = tmp_path / "cv.docx"
    p.write_bytes(docx_bytes(["مشرف أمن في مجمع تجاري بالخبر", "خبرة 7 سنوات"]))
    ex = extract(p, "docx", NoneOCR(), make_settings(), "none")
    assert ex.status == "ready" and ex.method == "docx"
    assert "مشرف أمن" in ex.text and "رخصة | سارية" in ex.text


def test_corrupt_and_encrypted(tmp_path: Path):
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"%PDF-1.4\n" + b"\x00garbage" * 100)
    ex = extract(bad, "pdf", NoneOCR(), make_settings(), "none")
    assert ex.status == "unreadable" and ex.error
    enc = tmp_path / "enc.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "secret cv")
    doc.save(str(enc), encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="o", user_pw="u")
    ex = extract(enc, "pdf", NoneOCR(), make_settings(), "none")
    assert ex.status == "unreadable" and "encrypted" in (ex.error or "")


@pytest.mark.skipif(not __import__("shutil").which("soffice"), reason="LibreOffice not installed")
def test_rtf_through_libreoffice(tmp_path: Path):
    p = tmp_path / "cv.rtf"
    p.write_bytes(
        b"{\\rtf1\\ansi Security guard with a valid licence and first aid training since 2019. Supervised a team of twelve guards.}"
    )
    ex = extract(p, "rtf", NoneOCR(), make_settings(), "none")
    assert ex.status == "ready" and ex.method.startswith("libreoffice") and "licence" in ex.text

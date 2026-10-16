"""Word version of the guide, same content model as the PDF.

Rules (Nadine, round 2): bidi in every paragraph style; no jc="right" together with bidi (start
alignment comes from bidi); real Word lists (numbering definitions); English runs without rtl;
bidi in sectPr before docGrid; an LTR «Code» style. The table of contents and the index use
bookmarks and PAGEREF fields, so Word (and LibreOffice) compute the page numbers themselves.
IBM Plex Sans Arabic is embedded (obfuscated TTF, as Word does) so the file looks the same on
a computer without the font."""

from __future__ import annotations

import copy
import os
import re
import uuid

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from model import Bullets, Chapter, Code, Diagram, EnLine, Field, H3, Img, Note, P, Page, Rule, Table, Tech
from render_html import LATIN_RUN

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
GUIDE = os.path.join(ROOT, "docs", "guide")
FONT = "IBM Plex Sans Arabic"
MONO = "IBM Plex Mono"
TEXT_W_CM = 17.0
BLUE = RGBColor(0x1E, 0x4D, 0xB5)
MUTED = RGBColor(0x56, 0x62, 0x7C)


def el(tag: str, **attrs) -> OxmlElement:
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


# child order of the OOXML property elements (Word refuses files with misordered children)
SEQ = {
    "w:rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath".split(),
    "w:pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange".split(),
    "w:tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription".split(),
    "w:tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark".split(),
}


def put(parent, child) -> None:
    """Insert `child` into a property element at its schema position (replacing the same tag)."""
    pname = parent.tag.split("}")[1]
    seq = SEQ["w:" + pname]
    cname = child.tag.split("}")[1]
    for old in parent.findall(child.tag):
        parent.remove(old)
    idx = seq.index(cname)
    for existing in parent:
        en = existing.tag.split("}")[1] if "}" in existing.tag else existing.tag
        if en in seq and seq.index(en) > idx:
            existing.addprevious(child)
            return
    parent.append(child)


SETTINGS_SEQ = ("writeProtection view zoom removePersonalInformation removeDateAndTime doNotDisplayPageBoundaries displayBackgroundShape "
                "printPostScriptOverText printFractionalCharacterWidth printFormsData embedTrueTypeFonts embedSystemFonts saveSubsetFonts saveFormsData "
                "mirrorMargins alignBordersAndEdges bordersDoNotSurroundHeader bordersDoNotSurroundFooter gutterAtTop hideSpellingErrors hideGrammaticalErrors "
                "activeWritingStyle proofState formsDesign attachedTemplate linkStyles stylePaneFormatFilter stylePaneSortMethod documentType mailMerge "
                "revisionView trackRevisions doNotTrackMoves doNotTrackFormatting documentProtection autoFormatOverride styleLockTheme styleLockQFSet "
                "defaultTabStop autoHyphenation consecutiveHyphenLimit hyphenationZone doNotHyphenateCaps showEnvelope summaryLength clickAndTypeStyle "
                "defaultTableStyle evenAndOddHeaders bookFoldRevPrinting bookFoldPrinting bookFoldPrintingSheets drawingGridHorizontalSpacing "
                "drawingGridVerticalSpacing displayHorizontalDrawingGridEvery displayVerticalDrawingGridEvery doNotUseMarginsForDrawingGridOrigin "
                "drawingGridHorizontalOrigin drawingGridVerticalOrigin doNotShadeFormData noPunctuationKerning characterSpacingControl printTwoOnOne "
                "strictFirstAndLastChars noLineBreaksAfter noLineBreaksBefore savePreviewPicture doNotValidateAgainstSchema saveInvalidXml ignoreMixedContent "
                "alwaysShowPlaceholderText doNotDemarcateInvalidXml saveXmlDataOnly useXSLTWhenSaving saveThroughXslt showXMLTags alwaysMergeEmptyNamespace "
                "updateFields hdrShapeDefaults footnotePr endnotePr compat docVars rsids mathPr attachedSchema themeFontLang clrSchemeMapping "
                "doNotIncludeSubdocsInStats doNotAutoCompressPictures forceUpgrade captions readModeInkLockDown smartTagType schemaLibrary shapeDefaults "
                "doNotEmbedSmartTags decimalSymbol listSeparator").split()


def put_settings(settings, child) -> None:
    name = child.tag.split("}")[1]
    for old in settings.findall(child.tag):
        settings.remove(old)
    idx = SETTINGS_SEQ.index(name)
    for existing in settings:
        en = existing.tag.split("}")[1] if "}" in existing.tag else existing.tag
        if en in SETTINGS_SEQ and SETTINGS_SEQ.index(en) > idx:
            existing.addprevious(child)
            return
    settings.append(child)


def set_bidi(pPr, on: bool = True) -> None:
    b = el("w:bidi")
    if not on:
        b.set(qn("w:val"), "0")
    put(pPr, b)


def style_fonts(style, font: str = FONT, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None) -> None:
    rPr = style.element.get_or_add_rPr()
    fonts = rPr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = el("w:rFonts")
        put(rPr, fonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), font)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if fonts.get(qn(a)) is not None:
            del fonts.attrib[qn(a)]
    if size:
        style.font.size = Pt(size)
        put(rPr, el("w:szCs", **{"w:val": int(size * 2)}))
    if bold is not None:
        style.font.bold = bold
        if bold:
            put(rPr, el("w:bCs"))
    if color is not None:
        style.font.color.rgb = color


class Docx:
    def __init__(self) -> None:
        self.doc = Document()
        self.bm = 0
        self._styles()
        self._section()
        self.list_num_abstract = self._abstract_of("List Number")

    # ------------------------------------------------------------------ setup
    def _styles(self) -> None:
        st = self.doc.styles
        for name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number", "Caption", "Title", "Subtitle", "Quote"):
            s = st[name]
            set_bidi(s.element.get_or_add_pPr(), True)
            style_fonts(s)
        style_fonts(st["Normal"], size=11)
        st["Normal"].paragraph_format.space_after = Pt(4)
        st["Normal"].paragraph_format.line_spacing = 1.25
        style_fonts(st["Heading 1"], size=22, bold=True, color=BLUE)
        style_fonts(st["Heading 2"], size=15, bold=True, color=RGBColor(0x1F, 0x2B, 0x47))
        style_fonts(st["Heading 3"], size=12.5, bold=True, color=BLUE)
        style_fonts(st["Caption"], size=9.5, bold=False, color=MUTED)
        st["Caption"].font.italic = False
        st["Heading 1"].paragraph_format.space_after = Pt(10)
        st["Heading 2"].paragraph_format.space_after = Pt(6)
        for name in ("Heading 1", "Heading 2", "Heading 3"):
            st[name].paragraph_format.keep_with_next = True
        # LTR styles: code boxes, the technical lines, English lines
        from docx.enum.style import WD_STYLE_TYPE

        for name, font, size, color in (("Code", MONO, 9, RGBColor(0x1F, 0x2B, 0x47)), ("Tech Line", MONO, 8.5, RGBColor(0x2C, 0x3A, 0x58)), ("English Line", FONT, 10.5, MUTED)):
            s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = st["Normal"]
            set_bidi(s.element.get_or_add_pPr(), False)
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            s.paragraph_format.space_after = Pt(1)
            s.paragraph_format.line_spacing = 1.1
            style_fonts(s, font=font, size=size, color=color)
            if font == MONO:  # an Arabic label inside a code line takes the Arabic font, not a fallback
                s.element.get_or_add_rPr().find(qn("w:rFonts")).set(qn("w:cs"), FONT)
        code_pPr = st["Code"].element.get_or_add_pPr()
        put(code_pPr, el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": "EEF1F7"}))

    def _section(self) -> None:
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Cm(1.6), Cm(1.8), Cm(2.0), Cm(2.0)
        sp = sec._sectPr
        grid = sp.find(qn("w:docGrid"))
        b = el("w:bidi")
        if grid is not None:
            grid.addprevious(b)
        else:
            sp.append(b)
        # page number in the footer
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._field(fp, "PAGE", "1")

    def _abstract_of(self, style_name: str) -> str | None:
        s = self.doc.styles[style_name].element
        num = s.find(qn("w:pPr") + "/" + qn("w:numPr") + "/" + qn("w:numId"))
        if num is None:
            return None
        numbering = self.doc.part.numbering_part.element
        for n in numbering.findall(qn("w:num")):
            if n.get(qn("w:numId")) == num.get(qn("w:val")):
                return n.find(qn("w:abstractNumId")).get(qn("w:val"))
        return None

    def _new_numbered_list(self) -> int:
        """A new w:num on the «List Number» definition that restarts at 1."""
        numbering = self.doc.part.numbering_part.element
        ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
        new_id = max(ids or [0]) + 1
        num = el("w:num", **{"w:numId": new_id})
        num.append(el("w:abstractNumId", **{"w:val": self.list_num_abstract}))
        ov = el("w:lvlOverride", **{"w:ilvl": 0})
        ov.append(el("w:startOverride", **{"w:val": 1}))
        num.append(ov)
        numbering.append(num)
        return new_id

    # ------------------------------------------------------------------ text
    def runs(self, p, text: str, bold: bool = False, color: RGBColor | None = None, size: float | None = None) -> None:
        """Arabic runs with rtl, Latin runs (English, numbers with letters, paths) without it."""
        pos = 0
        for m in LATIN_RUN.finditer(text or ""):
            if m.start() > pos:
                self._run(p, text[pos:m.start()], True, bold, color, size)
            self._run(p, m.group(0), False, bold, color, size)
            pos = m.end()
        if pos < len(text or ""):
            self._run(p, text[pos:], True, bold, color, size)

    def _run(self, p, t: str, rtl: bool, bold: bool, color, size) -> None:
        parts = t.split("\n")
        for i, part in enumerate(parts):
            if i:
                p.add_run().add_break(WD_BREAK.LINE)
            if not part:
                continue
            r = p.add_run(part)
            if bold:
                r.bold = True
                put(r._r.get_or_add_rPr(), el("w:bCs"))
            if color is not None:
                r.font.color.rgb = color
            if size:
                r.font.size = Pt(size)
                put(r._r.get_or_add_rPr(), el("w:szCs", **{"w:val": int(size * 2)}))
            if rtl:
                put(r._r.get_or_add_rPr(), el("w:rtl"))

    def _field(self, p, code: str, placeholder: str = "") -> None:
        r1 = p.add_run()
        r1._r.append(el("w:fldChar", **{"w:fldCharType": "begin"}))
        r2 = p.add_run()
        it = el("w:instrText")
        it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        it.text = f" {code} "
        r2._r.append(it)
        r3 = p.add_run()
        r3._r.append(el("w:fldChar", **{"w:fldCharType": "separate"}))
        p.add_run(placeholder)
        r5 = p.add_run()
        r5._r.append(el("w:fldChar", **{"w:fldCharType": "end"}))

    def bookmark(self, p, name: str) -> None:
        self.bm += 1
        p._p.insert(1 if p._p.pPr is not None else 0, el("w:bookmarkStart", **{"w:id": self.bm, "w:name": name}))
        p._p.append(el("w:bookmarkEnd", **{"w:id": self.bm}))

    def cached(self, pid: str) -> str:
        """The page shown before Word updates the fields (and by viewers that never do)."""
        return str(self.pages.get(self.bm_name(pid), "#"))

    @staticmethod
    def bm_name(pid: str) -> str:
        return "pg_" + re.sub(r"[^A-Za-z0-9_]", "_", pid)[:36]

    # ------------------------------------------------------------------ blocks
    def para(self, text: str, style: str = "Normal", **kw):
        p = self.doc.add_paragraph(style=style)
        self.runs(p, text, **kw)
        return p

    def heading(self, text: str, level: int, pid: str, page_break: bool) -> None:
        p = self.doc.add_paragraph(style=f"Heading {level}")
        p.paragraph_format.page_break_before = page_break
        self.runs(p, text)
        self.bookmark(p, self.bm_name(pid))

    def shaded_box(self, fill: str, lines: list[tuple[str, str]], border: str = "D9DFEB") -> None:
        """A one-cell table as a box: lines are (style, text)."""
        t = self.doc.add_table(rows=1, cols=1)
        tblPr = t._tbl.tblPr
        put(tblPr, el("w:bidiVisual"))
        borders = el("w:tblBorders")
        for side in ("top", "left", "bottom", "right"):
            borders.append(el(f"w:{side}", **{"w:val": "single", "w:sz": 6, "w:space": 0, "w:color": border}))
        put(tblPr, borders)
        cell = t.cell(0, 0)
        put(cell._tc.get_or_add_tcPr(), el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))
        first = True
        for style, text in lines:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.style = self.doc.styles[style]
            if style in ("Code", "Tech Line", "English Line"):
                self._run(p, text, False, False, None, None)
            else:
                self.runs(p, text, bold=style == "Normal-bold")
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def block(self, b) -> None:
        if isinstance(b, P):
            for line in b.text.split("\n"):
                self.para(line)
        elif isinstance(b, H3):
            p = self.doc.add_paragraph(style="Heading 3")
            self.runs(p, b.text)
        elif isinstance(b, Bullets):
            if b.numbered:
                nid = self._new_numbered_list()
            for it in b.items:
                p = self.doc.add_paragraph(style="List Number" if b.numbered else "List Bullet")
                if b.numbered:
                    numPr = p._p.get_or_add_pPr().get_or_add_numPr()
                    numPr.get_or_add_ilvl().val = 0
                    numPr.get_or_add_numId().val = nid
                self.runs(p, it)
        elif isinstance(b, Img):
            path = os.path.join(GUIDE, b.path)
            with Image.open(path) as im:
                w, h = im.size
            width = TEXT_W_CM * b.width
            height = width * h / w
            if height > b.max_h_cm:
                width, height = width * b.max_h_cm / height, b.max_h_cm
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(path, width=Cm(width))
            c = self.doc.add_paragraph(style="Caption")
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self.runs(c, b.caption)
        elif isinstance(b, Diagram):
            self.block(Img(b.png, b.caption, 0.86, 16))  # the same share of the width as in the PDF
        elif isinstance(b, Code):
            self.shaded_box("EEF1F7", [("Code", ln) for ln in b.text.split("\n")])
        elif isinstance(b, Tech):
            self.shaded_box("F1F3F7", [("Normal", b.title)] + [("Tech Line", ln) for ln in b.lines])
        elif isinstance(b, Rule):
            t = self.doc.add_table(rows=1, cols=1)
            put(t._tbl.tblPr, el("w:bidiVisual"))
            cell = t.cell(0, 0)
            put(cell._tc.get_or_add_tcPr(), el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": "1E4DB5"}))
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self.runs(p, b.text, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=14)
            self.doc.add_paragraph()
        elif isinstance(b, Note):
            fill = {"info": "E9EFFC", "warn": "FFF4E0", "ok": "E6F6EC"}[b.kind]
            self.shaded_box(fill, [("Normal", ln) for ln in b.text.split("\n")], border={"info": "1E4DB5", "warn": "8A5300", "ok": "166C3C"}[b.kind])
        elif isinstance(b, EnLine):
            p = self.doc.add_paragraph(style="English Line")
            self._run(p, b.text, False, False, None, None)
        elif isinstance(b, Field):
            p = self.doc.add_paragraph()
            self.runs(p, b.label + ": ", bold=True, color=BLUE)
            self.runs(p, b.text)
        elif isinstance(b, Table):
            self.table(b)
        else:
            raise TypeError(type(b))

    def table(self, b: Table) -> None:
        t = self.doc.add_table(rows=1, cols=len(b.head))
        t.style = self.doc.styles["Table Grid"]
        put(t._tbl.tblPr, el("w:bidiVisual"))
        size = 8 if b.small else 9.5
        for j, h in enumerate(b.head):
            c = t.rows[0].cells[j]
            put(c._tc.get_or_add_tcPr(), el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": "E9EFFC"}))
            self.runs(c.paragraphs[0], h, bold=True, size=size)
        trPr = t.rows[0]._tr.get_or_add_trPr()
        trPr.append(el("w:tblHeader"))
        for r in b.rows:
            cells = t.add_row().cells
            for j, v in enumerate(r):
                p = cells[j].paragraphs[0]
                if j in b.ltr_cols:
                    p.style = self.doc.styles["Tech Line"]
                    self._run(p, v, False, False, None, size - 0.5)
                else:
                    self.runs(p, v, size=size)
        if b.widths:
            # the grid columns too, and a fixed layout: LibreOffice reads the grid, Word the cells
            t.autofit = False
            total = TEXT_W_CM
            for j, w in enumerate(b.widths):
                t.columns[j].width = Cm(total * w)
            for row in t.rows:
                for j, w in enumerate(b.widths):
                    row.cells[j].width = Cm(total * w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ------------------------------------------------------------------ whole document
    def toc(self, chapters: list[Chapter]) -> None:
        """A borderless two-column table (title | page): the same in Word and LibreOffice."""
        rows = []
        for ch in chapters:
            if ch.id in ("cover", "toc"):
                continue
            rows.append((ch.title, ch.id, True))
            rows += [(pg.title, pg.id, False) for pg in ch.pages if pg.title and not pg.id.endswith("-intro")]
        t = self.doc.add_table(rows=0, cols=2)
        put(t._tbl.tblPr, el("w:bidiVisual"))
        for text, pid, top in rows:
            cells = t.add_row().cells
            cells[0].width, cells[1].width = Cm(TEXT_W_CM - 1.6), Cm(1.6)
            p = cells[0].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if not top:
                p.paragraph_format.first_line_indent = Cm(0.6)
            self.runs(p, text, bold=top, size=10.5 if top else 9.5)
            q = cells[1].paragraphs[0]
            q.paragraph_format.space_after = Pt(0)
            self._field(q, f"PAGEREF {self.bm_name(pid)} \\h", self.cached(pid))

    def index(self, entries: list[tuple[str, list[str]]]) -> None:
        last = ""
        for term, pids in entries:
            letter = re.sub(r"^ال", "", term)[:1]
            if letter != last:
                last = letter
                self.para(last, bold=True, color=BLUE)
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            self.runs(p, term + " ", size=9.5)
            for i, pid in enumerate(pids):
                if i:
                    p.add_run("، ")
                self._field(p, f"PAGEREF {self.bm_name(pid)} \\h", self.cached(pid))

    def embed_fonts(self, ttf: dict[str, str]) -> None:
        """Embed TTF files the way Word does (obfuscated with the font key)."""
        fpart = self.doc.part.part_related_by(RT.FONT_TABLE) if any(r.reltype == RT.FONT_TABLE for r in self.doc.part.rels.values()) else None
        if fpart is None:
            return
        from lxml import etree

        root = etree.fromstring(fpart.blob)
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        for name in (FONT, MONO):
            f = root.find(f"w:font[@w:name='{name}']", ns)
            if f is None:
                f = etree.SubElement(root, qn("w:font"))
                f.set(qn("w:name"), name)
            for kind, path in ttf.items():
                fam, style = kind.split(":")
                if fam != name:
                    continue
                key = "{" + str(uuid.uuid4()).upper() + "}"
                data = bytearray(open(path, "rb").read())
                k = bytes.fromhex(key.strip("{}").replace("-", ""))[::-1]
                for i in range(32):
                    data[i] ^= k[i % 16]
                part = Part(PackURI(f"/word/fonts/{re.sub(r'[^A-Za-z]', '', name)}{style}.odttf"), "application/vnd.openxmlformats-officedocument.obfuscatedFont", bytes(data), fpart.package)
                rid = fpart.relate_to(part, RT.FONT)
                e = etree.SubElement(f, qn({"Regular": "w:embedRegular", "Bold": "w:embedBold"}[style]))
                e.set(qn("r:id"), rid)
                e.set(qn("w:fontKey"), key)
        fpart._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        settings = self.doc.settings.element
        put_settings(settings, el("w:embedTrueTypeFonts"))
        put_settings(settings, el("w:updateFields", **{"w:val": "true"}))


def render(chapters: list[Chapter], index_entries: list[tuple[str, list[str]]], cover: dict, out: str, ttf: dict[str, str] | None = None,
           pages: dict[str, int] | None = None) -> None:
    """`pages` (bookmark → page, read from a first conversion) fills the PAGEREF fields' shown value."""
    d = Docx()
    d.pages = pages or {}
    p = d.doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Cm(7)
    d.runs(p, cover["title"])
    p = d.doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.runs(p, cover["sub"])
    p = d.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Cm(4)
    d.runs(p, cover["meta"], color=MUTED)
    for ch in chapters:
        d.heading(ch.title, 1, ch.id, True)
        if ch.id == "toc":
            d.toc(chapters)
            continue
        if ch.id == "index":
            d.index(index_entries)
            continue
        if ch.intro:
            d.para(ch.intro, color=MUTED)
        for j, pg in enumerate(ch.pages):
            if pg.title:
                d.heading(pg.title, 2, pg.id, pg.break_before and j > 0)
            for b in pg.blocks:
                d.block(b)
    if ttf:
        d.embed_fonts(ttf)
    d.doc.save(out)
